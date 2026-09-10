"""
Compliance Module — Read-Only Sheets Client
============================================
Fetches and parses the Compliance Management spreadsheet in 100% read-only mode.
Because the document on Google Drive is an Excel (.xlsx) file, this client
downloads the binary stream in-memory via Google Drive API and parses the
GEBB1, GEBB2, and GETT worksheets using openpyxl.

Strictly READ-ONLY: Never writes or alters anything on Google Drive or Sheets.
"""

import io
import logging
from datetime import date, datetime
from typing import Any, Dict, List, Optional

import openpyxl
from dateutil import parser as date_parser
# pyrefly: ignore [missing-import]
from google.auth.transport.requests import AuthorizedSession
# pyrefly: ignore [missing-import]
from google.oauth2.service_account import Credentials

from compliance.config import (
    BUILDING_NAME_MAPPING,
    COL_APPLICABILITY,
    COL_CATEGORY,
    COL_CONTROL_GROUP,
    COL_DUE_DATE,
    COL_EVIDENCE,
    COL_ID,
    COL_ISSUE_DATE,
    COL_OWNER,
    COL_REMARKS,
    COL_REQUIREMENT,
    COL_STATUS,
    COMPLIANCE_BUILDING_TABS,
    COMPLIANCE_SA_EMAIL,
    COMPLIANCE_SA_PRIVATE_KEY,
    COMPLIANCE_SHEET_ID,
    DATA_START_ROW_INDEX,
    HEADER_ROW_INDEX,
)

logger = logging.getLogger(__name__)

SCOPES = ["https://www.googleapis.com/auth/drive.readonly"]


def _get_authorized_session() -> AuthorizedSession:
    """Create an authorized Google HTTP session using existing service account credentials."""
    if not COMPLIANCE_SA_EMAIL or not COMPLIANCE_SA_PRIVATE_KEY:
        raise RuntimeError(
            "Service account credentials not configured for Compliance module. "
            "Check FACILITIES_SERVICE_ACCOUNT_EMAIL and FACILITIES_PRIVATE_KEY."
        )

    pk = COMPLIANCE_SA_PRIVATE_KEY.strip().strip("'").strip('"').replace("\\n", "\n")

    # Derive project_id from service account domain if available
    project_id = "facilities-tracker"
    if COMPLIANCE_SA_EMAIL and "@" in COMPLIANCE_SA_EMAIL:
        domain = COMPLIANCE_SA_EMAIL.split("@")[1]
        if "." in domain:
            project_id = domain.split(".")[0]

    creds_info = {
        "type": "service_account",
        "project_id": project_id,
        "private_key": pk,
        "client_email": COMPLIANCE_SA_EMAIL,
        "token_uri": "https://oauth2.googleapis.com/token",
    }

    creds = Credentials.from_service_account_info(creds_info, scopes=SCOPES)
    return AuthorizedSession(creds)


def parse_cell_date(value: Any) -> Optional[date]:
    """
    Safely parse cell value into a Python date object.
    Supports datetime.datetime, datetime.date, and formatted date strings (e.g. '21-Oct-2023', '2026-10-20').
    """
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    val_str = str(value).strip()
    if not val_str or val_str.lower() in ("none", "n/a", "-", "null", "nil", ""):
        return None

    try:
        # Handles dayfirst formats like 21-Oct-2023 or 20/10/2026
        dt = date_parser.parse(val_str, dayfirst=True)
        return dt.date()
    except (ValueError, OverflowError):
        logger.debug(f"Could not parse date string: {val_str}")
        return None


def fetch_compliance_sheet_stream(sheet_id: Optional[str] = None) -> bytes:
    """
    Download the binary stream of the .xlsx file from Google Drive via read-only Drive API.
    Does NOT write anything to Google Drive.
    """
    target_id = sheet_id or COMPLIANCE_SHEET_ID
    if not target_id:
        raise ValueError("COMPLIANCE_SHEET_ID is not configured.")

    session = _get_authorized_session()
    url = f"https://www.googleapis.com/drive/v3/files/{target_id}?alt=media"

    logger.info(f"Downloading compliance spreadsheet (ID: {target_id}) via Drive API...")
    response = session.get(url, timeout=45)

    if response.status_code != 200:
        logger.error(f"Failed to download compliance sheet: {response.status_code} - {response.text[:300]}")
        response.raise_for_status()

    content = response.content
    logger.info(f"Successfully downloaded compliance workbook ({len(content)} bytes).")
    return content


def parse_compliance_records_from_stream(
    xlsx_bytes: bytes,
    tabs_to_read: Optional[List[str]] = None,
) -> List[Dict[str, Any]]:
    """
    Parse compliance items from GEBB1, GEBB2, and GETT worksheets in the provided workbook bytes.
    Management Dashboard is explicitly skipped.
    """
    tabs = tabs_to_read or COMPLIANCE_BUILDING_TABS
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), data_only=True)

    records: List[Dict[str, Any]] = []

    for tab_name in tabs:
        if tab_name not in wb.sheetnames:
            logger.warning(f"Tab '{tab_name}' not found in compliance workbook. Available: {wb.sheetnames}")
            continue

        sheet = wb[tab_name]
        logger.info(f"Reading compliance tab '{tab_name}' (rows: {sheet.max_row})...")

        # Discover dynamic column positions from row 4 if present
        col_map = {
            "id": COL_ID,
            "control_group": COL_CONTROL_GROUP,
            "category": COL_CATEGORY,
            "requirement": COL_REQUIREMENT,
            "applicability": COL_APPLICABILITY,
            "issue_date": COL_ISSUE_DATE,
            "due_date": COL_DUE_DATE,
            "status": COL_STATUS,
            "evidence": COL_EVIDENCE,
            "remarks": COL_REMARKS,
            "owner": COL_OWNER,
        }

        for col_idx in range(1, min(sheet.max_column + 1, 30)):
            hdr_val = sheet.cell(row=HEADER_ROW_INDEX, column=col_idx).value
            if not hdr_val:
                continue
            hdr_lower = str(hdr_val).strip().lower()
            if hdr_lower == "id":
                col_map["id"] = col_idx
            elif "control group" in hdr_lower:
                col_map["control_group"] = col_idx
            elif "category" in hdr_lower:
                col_map["category"] = col_idx
            elif "evidence" in hdr_lower or "ref" in hdr_lower:
                col_map["evidence"] = col_idx
            elif "requirement" in hdr_lower:
                col_map["requirement"] = col_idx
            elif "applicab" in hdr_lower:
                col_map["applicability"] = col_idx
            elif "issue date" in hdr_lower or "last done" in hdr_lower:
                col_map["issue_date"] = col_idx
            elif "valid till" in hdr_lower or "next due" in hdr_lower:
                col_map["due_date"] = col_idx
            elif "status" in hdr_lower:
                col_map["status"] = col_idx
            elif "action" in hdr_lower or "remark" in hdr_lower:
                col_map["remarks"] = col_idx
            elif "owner" in hdr_lower:
                col_map["owner"] = col_idx

        # Iterate over data rows
        for r_idx in range(DATA_START_ROW_INDEX, sheet.max_row + 1):
            raw_id = sheet.cell(row=r_idx, column=col_map["id"]).value
            if not raw_id:
                continue

            compliance_id = str(raw_id).strip()
            if not compliance_id or compliance_id.lower() == "id":
                continue

            control_group = sheet.cell(row=r_idx, column=col_map["control_group"]).value or ""
            category = sheet.cell(row=r_idx, column=col_map["category"]).value or ""
            requirement = sheet.cell(row=r_idx, column=col_map["requirement"]).value or ""
            applicability = sheet.cell(row=r_idx, column=col_map["applicability"]).value or ""
            raw_issue_date = sheet.cell(row=r_idx, column=col_map["issue_date"]).value
            raw_due_date = sheet.cell(row=r_idx, column=col_map["due_date"]).value
            status = sheet.cell(row=r_idx, column=col_map["status"]).value or ""
            evidence = sheet.cell(row=r_idx, column=col_map["evidence"]).value or ""
            remarks = sheet.cell(row=r_idx, column=col_map["remarks"]).value or ""
            owner = sheet.cell(row=r_idx, column=col_map["owner"]).value or ""

            issue_date = parse_cell_date(raw_issue_date)
            due_date = parse_cell_date(raw_due_date)

            building_display = BUILDING_NAME_MAPPING.get(tab_name, tab_name)

            records.append({
                "building": tab_name,
                "building_name": building_display,
                "compliance_id": compliance_id,
                "control_group": str(control_group).strip(),
                "category": str(category).strip(),
                "requirement": str(requirement).strip(),
                "applicability": str(applicability).strip(),
                "issue_date": issue_date,
                "issue_date_raw": str(raw_issue_date) if raw_issue_date is not None else "",
                "due_date": due_date,
                "due_date_raw": str(raw_due_date) if raw_due_date is not None else "",
                "status": str(status).strip(),
                "evidence": str(evidence).strip(),
                "remarks": str(remarks).strip(),
                "owner": str(owner).strip(),
                "row_index": r_idx,
            })

    logger.info(f"Parsed {len(records)} compliance items across {tabs}.")
    return records


def fetch_all_compliance_records(sheet_id: Optional[str] = None) -> List[Dict[str, Any]]:
    """Download workbook and parse all valid records from GEBB1, GEBB2, and GETT."""
    content = fetch_compliance_sheet_stream(sheet_id=sheet_id)
    return parse_compliance_records_from_stream(content)
