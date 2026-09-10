"""
Unit & Integration Tests for Compliance Reminder Module
========================================================
Tests cover:
  1. Date parsing utilities
  2. In-memory .xlsx workbook parsing (ignoring Management Dashboard)
  3. Meta WhatsApp Utility Template payload construction (11 body parameters)
  4. WhatsApp template dispatcher to Anoop Sir
  5. Duplicate prevention and audit logging
  6. Daily scheduler job orchestration and date filtering
"""

import io
from datetime import date, datetime
from unittest.mock import MagicMock, patch

import openpyxl
# pyrefly: ignore [missing-import]
import pytest

from compliance.config import BUILDING_NAME_MAPPING, COMPLIANCE_TEMPLATE_NAME
from compliance.notifier import (
    build_compliance_template_payload,
    format_display_date,
    resolve_anoop_phone_number,
    send_compliance_reminder_to_anoop,
)
from compliance.scheduler_job import run_compliance_daily_check
from compliance.sheets_client import (
    parse_cell_date,
    parse_compliance_records_from_stream,
)
from compliance.tracker import (
    has_reminder_been_sent_today,
    log_reminder_attempt,
)


# ── 1. Date Parsing Tests ───────────────────────────────────────────────────

def test_parse_cell_date():
    # From datetime
    dt = datetime(2026, 10, 20, 14, 30)
    assert parse_cell_date(dt) == date(2026, 10, 20)

    # From date
    d = date(2026, 10, 20)
    assert parse_cell_date(d) == date(2026, 10, 20)

    # From string formats
    assert parse_cell_date("20-Oct-2026") == date(2026, 10, 20)
    assert parse_cell_date("21-Jun-2023") == date(2023, 6, 21)
    assert parse_cell_date("2026-10-20") == date(2026, 10, 20)

    # Invalid / empty
    assert parse_cell_date(None) is None
    assert parse_cell_date("") is None
    assert parse_cell_date("-") is None
    assert parse_cell_date("N/A") is None


def test_format_display_date():
    assert format_display_date(date(2026, 10, 20)) == "20-Oct-2026"
    assert format_display_date(datetime(2023, 6, 21, 0, 0)) == "21-Jun-2023"
    assert format_display_date(None) == "N/A"
    assert format_display_date("20-Oct-2026") == "20-Oct-2026"


# ── 2. Workbook Stream Parsing Tests ────────────────────────────────────────

def _create_test_xlsx_bytes() -> bytes:
    """Create a synthetic in-memory workbook mimicking the real compliance file."""
    wb = openpyxl.Workbook()

    # Management Dashboard (should be ignored)
    ws_dash = wb.active
    ws_dash.title = "Management Dashboard"
    ws_dash.cell(row=1, column=1, value="Dashboard Summary")

    # Building tabs: GEBB1, GEBB2, GETT
    headers = [
        "ID", "Control Group", "Category", "Requirement / Document",
        "Authority / Source", "Frequency / Trigger", "Classification",
        "Applicability", "Issue Date / Last Done", "Valid Till / Next Due",
        "Status", "Evidence / Document Ref", "Action / Remarks", "Owner",
    ]

    for tab_name in ["GEBB1", "GEBB2", "GETT"]:
        ws = wb.create_sheet(title=tab_name)
        # Row 4 is header
        for col_idx, h in enumerate(headers, 1):
            ws.cell(row=4, column=col_idx, value=h)

        # Row 5 data
        ws.cell(row=5, column=1, value="RG-01")
        ws.cell(row=5, column=2, value="Renewable")
        ws.cell(row=5, column=3, value="Fire & Life Safety")
        ws.cell(row=5, column=4, value="Fire NOC / Fire Safety Certificate")
        ws.cell(row=5, column=8, value="Yes")
        ws.cell(row=5, column=9, value=datetime(2023, 10, 21))
        ws.cell(row=5, column=10, value=datetime(2026, 10, 20))
        ws.cell(row=5, column=11, value="Due ≤60 Days")
        ws.cell(row=5, column=12, value="DOC-FIRE-01")
        ws.cell(row=5, column=13, value="Initiate renewal process.")
        ws.cell(row=5, column=14, value="Anoop Sir")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_compliance_records_from_stream():
    content = _create_test_xlsx_bytes()
    records = parse_compliance_records_from_stream(content)

    # 3 tabs * 1 record = 3 records
    assert len(records) == 3

    buildings = [r["building"] for r in records]
    assert "GEBB1" in buildings
    assert "GEBB2" in buildings
    assert "GETT" in buildings
    assert "Management Dashboard" not in buildings

    gebb1_rec = next(r for r in records if r["building"] == "GEBB1")
    assert gebb1_rec["compliance_id"] == "RG-01"
    assert gebb1_rec["building_name"] == "Bay 1"
    assert gebb1_rec["control_group"] == "Renewable"
    assert gebb1_rec["category"] == "Fire & Life Safety"
    assert gebb1_rec["requirement"] == "Fire NOC / Fire Safety Certificate"
    assert gebb1_rec["due_date"] == date(2026, 10, 20)
    assert gebb1_rec["status"] == "Due ≤60 Days"
    assert gebb1_rec["owner"] == "Anoop Sir"
    assert gebb1_rec["evidence"] == "DOC-FIRE-01"


# ── 3. WhatsApp Template Payload Builder Tests ──────────────────────────────

def test_build_compliance_template_payload():
    record = {
        "building": "GEBB1",
        "building_name": "Bay 1",
        "compliance_id": "RG-01",
        "control_group": "Renewable",
        "category": "Fire & Life Safety",
        "requirement": "Fire NOC / Fire Safety Certificate",
        "applicability": "Yes",
        "issue_date": date(2023, 10, 21),
        "due_date": date(2026, 10, 20),
        "status": "Due ≤60 Days",
        "owner": "Anoop Sir",
        "remarks": "Initiate renewal process.",
        "evidence": "NOC-1234",
    }

    payload = build_compliance_template_payload("919211501013", record)

    assert payload["messaging_product"] == "whatsapp"
    assert payload["to"] == "919211501013"
    assert payload["type"] == "template"
    assert payload["template"]["name"] == COMPLIANCE_TEMPLATE_NAME
    assert payload["template"]["language"]["code"] == "en"

    body_comp = payload["template"]["components"][0]
    assert body_comp["type"] == "body"

    params = [p["text"] for p in body_comp["parameters"]]
    assert len(params) == 11

    assert params[0] == "Bay 1"
    assert params[1] == "RG-01"
    assert params[2] == "Renewable"
    assert params[3] == "Fire & Life Safety"
    assert params[4] == "Fire NOC / Fire Safety Certificate"
    assert params[5] == "Yes"
    assert params[6] == "21-Oct-2023"
    assert params[7] == "20-Oct-2026"
    assert params[8] == "Due ≤60 Days"
    assert params[9] == "Anoop Sir"
    assert "Initiate renewal process." in params[10]
    assert "Ref: NOC-1234" in params[10]


def test_build_compliance_template_payload_fallback_empty_fields():
    record = {
        "building": "GETT",
        "building_name": "Trade Tower",
        "compliance_id": "RG-10",
        "control_group": "",
        "category": "",
        "requirement": "Lift inspection",
        "applicability": "",
        "issue_date": None,
        "due_date": date(2026, 10, 20),
        "status": "",
        "owner": "",
        "remarks": "",
        "evidence": "",
    }

    payload = build_compliance_template_payload("919211501013", record)
    params = [p["text"] for p in payload["template"]["components"][0]["parameters"]]

    assert len(params) == 11
    # Check fallback values
    for p in params:
        assert p is not None
        assert p != ""  # Meta rejects empty strings
    assert params[9] == "Anoop Sir"  # Default owner
    assert "Please review" in params[10] or "Initiate" in params[10]


# ── 4. WhatsApp Dispatcher Tests ─────────────────────────────────────────────

@patch("compliance.notifier.requests.post")
def test_send_compliance_reminder_to_anoop(mock_post):
    mock_resp = MagicMock()
    mock_resp.status_code = 200
    mock_resp.text = '{"messages":[{"id":"wamid.test"}]}'
    mock_resp.json.return_value = {"messages": [{"id": "wamid.test"}]}
    mock_post.return_value = mock_resp

    record = {
        "building": "GEBB2",
        "building_name": "Bay 2",
        "compliance_id": "RG-05",
        "control_group": "Renewable",
        "category": "Electrical",
        "requirement": "Transformer inspection",
        "applicability": "Yes",
        "issue_date": date(2025, 1, 1),
        "due_date": date(2026, 10, 20),
        "status": "Due Today",
        "owner": "Anoop Sir",
        "remarks": "Check oil levels",
        "evidence": "",
    }

    success, data, err = send_compliance_reminder_to_anoop(record, override_phone="919211501013")
    assert success is True
    assert err is None
    assert mock_post.called

    sent_payload = mock_post.call_args[1]["json"]
    assert sent_payload["to"] == "919211501013"
    assert sent_payload["template"]["name"] == "compliance_due_reminder"


# ── 5. Duplicate Prevention Tests ────────────────────────────────────────────

def test_duplicate_prevention_sqlite():
    import uuid
    test_cid = f"TEST-DEDUP-{uuid.uuid4().hex[:8].upper()}"
    test_building = "GEBB1"
    test_due = date(2026, 10, 20)
    test_sent_date = date(2026, 10, 20)

    # Before logging, it should not be marked as sent
    assert has_reminder_been_sent_today(test_cid, test_building, test_due, check_date=test_sent_date) is False


    # Log reminder
    log_reminder_attempt(
        compliance_id=test_cid,
        building=test_building,
        due_date=test_due,
        recipient_phone="919211501013",
        whatsapp_status="accepted",
        delivery_status="sent",
        sent_date=test_sent_date,
    )

    # Now it should report as already sent today
    assert has_reminder_been_sent_today(test_cid, test_building, test_due, check_date=test_sent_date) is True

    # On a different sent date, it should allow sending again
    another_date = date(2026, 10, 21)
    assert has_reminder_been_sent_today(test_cid, test_building, test_due, check_date=another_date) is False


# ── 6. Scheduled Job Orchestration & Filtering Tests ─────────────────────────

@patch("compliance.scheduler_job.send_compliance_reminder_to_anoop")
def test_run_compliance_daily_check_filtering(mock_send):
    mock_send.return_value = (True, {"id": "wamid.123"}, None)

    today = date(2026, 9, 10)
    mock_records = [
        # Item 1: Due TODAY -> Should send
        {
            "building": "GEBB1",
            "building_name": "Bay 1",
            "compliance_id": "TEST-DUE-TODAY",
            "due_date": today,
            "requirement": "Fire Safety Certificate",
        },
        # Item 2: Due TOMORROW -> Should NOT send
        {
            "building": "GEBB2",
            "building_name": "Bay 2",
            "compliance_id": "TEST-DUE-TOMORROW",
            "due_date": date(2026, 9, 11),
            "requirement": "Lift Inspection",
        },
        # Item 3: Due YESTERDAY -> Should NOT send
        {
            "building": "GETT",
            "building_name": "Trade Tower",
            "compliance_id": "TEST-DUE-YESTERDAY",
            "due_date": date(2026, 9, 9),
            "requirement": "Consent to Operate",
        },
    ]

    summary = run_compliance_daily_check(
        target_date=today,
        force_send=True,
        records_cache=mock_records,
    )

    assert summary["total_checked"] == 3
    assert summary["due_count"] == 1
    assert summary["sent_count"] == 1
    assert summary["skipped_duplicate"] == 0
    assert len(summary["due_items"]) == 1
    assert summary["due_items"][0]["compliance_id"] == "TEST-DUE-TODAY"

    mock_send.assert_called_once()
